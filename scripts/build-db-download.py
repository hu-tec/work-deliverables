"""Build a downloadable DB archive and manifest from reviewed local files."""

import csv
import json
import zipfile
from pathlib import Path
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DB = ROOT / "추가_산출물/DB_정제_설명문서"
ARCHIVE = ROOT / "downloads/DB_정제_자료.zip"
MANIFEST = ROOT / "db-files.js"


def main():
    files = sorted(
        file for file in DB.rglob("*")
        if file.is_file() and file.suffix.lower() in {".csv", ".md"}
    )
    files.extend(sorted((ROOT / "추가_산출물/DB_원본_엑셀").glob("*.xlsx")))
    with zipfile.ZipFile(ARCHIVE, "w", zipfile.ZIP_DEFLATED) as output:
        for file in files:
            output.write(file, file.relative_to(ROOT).as_posix())
    payload = []
    for file in files:
        item = {
            "path": file.relative_to(ROOT).as_posix(),
            "group": file.relative_to(DB).parent.as_posix() if file.is_relative_to(DB) else "원본 엑셀",
            "name": file.name,
            "size": file.stat().st_size,
            "type": file.suffix.lower().lstrip("."),
        }
        if item["type"] == "csv":
            with file.open(encoding="utf-8-sig", newline="") as source:
                item["rows"] = list(csv.reader(source))
        elif item["type"] == "md":
            item["content"] = file.read_text(encoding="utf-8", errors="replace")
        else:
            workbook = load_workbook(file, read_only=True, data_only=True)
            item["sheets"] = [
                {"name": sheet.title, "rows": [[str(value) if value is not None else "" for value in row]
                                               for row in sheet.iter_rows(values_only=True)]}
                for sheet in workbook
            ]
            workbook.close()
        payload.append(item)
    MANIFEST.write_text(
        "window.DB_FILES = " + json.dumps(payload, ensure_ascii=False).replace("</", "<\\/") + ";\n",
        encoding="utf-8",
    )
    print(f"DB download: {sum(file.suffix.lower() == '.csv' for file in files)} CSV, "
          f"{sum(file.suffix.lower() == '.md' for file in files)} Markdown, "
          f"{sum(file.suffix.lower() == '.xlsx' for file in files)} Excel")


if __name__ == "__main__":
    main()
